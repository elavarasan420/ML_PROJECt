from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

OUTPUT_PATH = r"c:\Users\ezhil\OneDrive\Desktop\AIML\disease_prediction_app\Disease_Prediction_Project_PPT.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(18, 38, 71)
BLUE = RGBColor(27, 122, 197)
LIGHT = RGBColor(240, 245, 250)
TEXT = RGBColor(35, 35, 35)
GREEN = RGBColor(39, 174, 96)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(100, 100, 100)


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.0), Inches(0.7))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.LEFT

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.6), Inches(0.9), Inches(12.0), Inches(0.4))
        sub_tf = sub.text_frame
        p2 = sub_tf.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(12)
        p2.font.color.rgb = RGBColor(75, 86, 96)


def add_bullets(slide, items, x, y, w, h, font_size=18, color=TEXT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 6
    tf.margin_right = 6
    tf.margin_top = 4
    tf.margin_bottom = 4

    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(6)


# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = LIGHT

band = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.7))
band.fill.solid()
band.fill.fore_color.rgb = NAVY
band.line.fill.background()

box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(10.5), Inches(0.8))
tf = box.text_frame
p = tf.paragraphs[0]
p.text = 'MediCare AI - Disease Prediction System'
p.font.size = Pt(28)
p.font.bold = True
p.font.color.rgb = NAVY
p.alignment = PP_ALIGN.LEFT

sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.0), Inches(0.6))
sub_tf = sub.text_frame
p2 = sub_tf.paragraphs[0]
p2.text = 'Machine Learning Based Cardiovascular Disease Risk Assessment using Flask Web Application'
p2.font.size = Pt(16)
p2.font.color.rgb = BLUE

cards = [
    ('Project Type', 'AI + Web App'),
    ('Purpose', 'Risk prediction and guidance'),
    ('Tech Stack', 'Python, Flask, ML'),
    ('Output', 'Prediction + history')
]
for idx, (title, value) in enumerate(cards):
    x = Inches(0.8 + (idx % 2) * 5.6)
    y = Inches(3.0 + (idx // 2) * 1.5)
    shape = slide.shapes.add_shape(1, x, y, Inches(4.7), Inches(1.1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = BLUE
    tfc = shape.text_frame
    tfc.word_wrap = True
    p1 = tfc.paragraphs[0]
    p1.text = title
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = BLUE
    p2 = tfc.add_paragraph()
    p2.text = value
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT

footer = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(8.2), Inches(0.4))
footer_tf = footer.text_frame
footer_p = footer_tf.paragraphs[0]
footer_p.text = 'Prepared for project presentation | Disease Prediction App'
footer_p.font.size = Pt(10)
footer_p.font.color.rgb = GRAY

# Slide 2: Overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
add_title(slide, '1. Project Overview', 'Introduction and vision')
add_bullets(slide, [
    'The project is a disease prediction system designed to estimate cardiovascular risk using clinical parameters.',
    'It combines machine learning and a user-friendly web interface to support early health awareness and risk assessment.',
    'The solution includes user authentication, prediction history, recovery suggestions, and downloadable reports.',
    'The goal is to provide a practical demonstration of AI-driven healthcare decision support with transparent outputs.'
], Inches(0.8), Inches(1.5), Inches(11.6), Inches(4.3), font_size=20)

# Slide 3: Problem Statement
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
add_title(slide, '2. Problem Statement', 'Why this project matters')
add_bullets(slide, [
    'Many users do not have ready access to timely medical evaluation for early cardiovascular symptoms.',
    'Healthcare decisions often require manual review of several parameters such as age, blood pressure, cholesterol, and symptoms.',
    'Late detection can increase the risk of severe complications and poor health outcomes.',
    'There is a need for a low-cost, explainable, and accessible digital system that supports preventive health monitoring.'
], Inches(0.8), Inches(1.5), Inches(11.6), Inches(4.3), font_size=20)

# Slide 4: Objectives
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
add_title(slide, '3. Objectives', 'Project goals')
add_bullets(slide, [
    'To build a disease prediction model that classifies risk as Low Risk or High Risk using clinical data.',
    'To design a Flask-based web application with secure registration, login, prediction, and user history support.',
    'To provide clear recovery suggestions and symptom-based recommendations in a user-friendly format.',
    'To create a complete end-to-end workflow from data input to prediction output and downloadable report generation.',
    'To demonstrate the real-world use of AI in healthcare through a practical project implementation.'
], Inches(0.8), Inches(1.5), Inches(11.6), Inches(4.8), font_size=18)

# Slide 5: System Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
add_title(slide, '4. System Architecture', 'Core workflow')
modules = [
    ('Input Layer', 'User enters health metrics and symptom data.'),
    ('ML Layer', 'Model predicts risk using normalized feature values.'),
    ('Web Layer', 'Flask app handles pages, routes, and APIs.'),
    ('Database Layer', 'SQLite stores users and prediction records.'),
    ('Output Layer', 'Prediction results and suggestions are displayed and saved.')
]
for idx, (title, desc) in enumerate(modules):
    x = Inches(0.7 + (idx % 2) * 6.0)
    y = Inches(1.6 + (idx // 2) * 1.7)
    shape = slide.shapes.add_shape(1, x, y, Inches(5.2), Inches(1.3))
    shape.fill.solid(); shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = BLUE
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT

# Slide 6: Features
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
add_title(slide, '5. Data and Features', 'Input parameters used for prediction')
feature_list = [
    'Age', 'Sex', 'Chest pain type', 'Resting blood pressure', 'Cholesterol',
    'Fasting blood sugar', 'Resting ECG', 'Max heart rate', 'Exercise angina',
    'Oldpeak', 'Slope', 'Major vessels', 'Thalassemia status', 'Optional symptom'
]
add_bullets(slide, feature_list[:7], Inches(0.9), Inches(1.6), Inches(5.2), Inches(4.8), font_size=18)
add_bullets(slide, feature_list[7:], Inches(6.8), Inches(1.6), Inches(5.4), Inches(4.8), font_size=18)

# Slide 7: Model Workflow
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
add_title(slide, '6. Model and Prediction Workflow', 'How the system works')
add_bullets(slide, [
    'The application loads a trained model and scaler from saved files.',
    'Clinical input values are converted to a DataFrame and normalized using the scaler.',
    'The machine learning model predicts whether the user is at Low Risk or High Risk.',
    'Prediction confidence is computed as the probability of the selected class.',
    'The app adds symptom-based insights and useful recovery recommendations.',
    'Results are stored in SQLite and can be reviewed from the prediction history section.'
], Inches(0.8), Inches(1.7), Inches(11.6), Inches(4.7), font_size=19)

# Slide 8: Features of app
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
add_title(slide, '7. Application Features', 'What the web app offers')
feature_cards = [
    ('User Authentication', 'Registration, login, logout, and session management.'),
    ('Prediction Page', 'Submit patient data and optional symptom details.'),
    ('Dashboard', 'Review recent results and health insights.'),
    ('History Tracking', 'Save past predictions and retrieve prior records.'),
    ('PDF Export', 'Download a PDF report of prediction history.'),
    ('Image Analysis', 'Upload an image and receive basic injury guidance.')
]
for idx, (title, desc) in enumerate(feature_cards):
    x = Inches(0.55 + (idx % 3) * 4.05)
    y = Inches(1.7 + (idx // 3) * 1.9)
    shape = slide.shapes.add_shape(1, x, y, Inches(3.55), Inches(1.4))
    shape.fill.solid(); shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = GREEN
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = GREEN
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(9)
    p2.font.color.rgb = TEXT

# Slide 9: Tech Stack
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
add_title(slide, '8. Technology Stack', 'Tools and frameworks used')
add_bullets(slide, [
    'Python for application logic and machine learning operations.',
    'Flask for the backend, routes, templates, and API development.',
    'Scikit-learn for model training and preprocessing.',
    'Pandas, NumPy, and Joblib for data handling and model persistence.',
    'SQLite for user account and prediction storage.',
    'HTML, CSS, and JavaScript for the frontend interface.',
    'Docker and Gunicorn for deployment-ready support.'
], Inches(0.8), Inches(1.6), Inches(11.6), Inches(4.8), font_size=18)

# Slide 10: Evaluation and Future Scope
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
add_title(slide, '9. Advantages, Limitations, and Future Scope', 'Project evaluation')
add_bullets(slide, [
    'Advantages: Easy to use, explainable, secure login, prediction history, and downloadable reports.',
    'Limitations: This is a demo application and should not replace clinical diagnosis or medical advice.',
    'Future scope: Add real hospital datasets, advanced models, analytics dashboards, and doctor-facing interfaces.',
    'Potential extension: Integrate with electronic health records, remote monitoring, and more advanced AI-based diagnostics.'
], Inches(0.8), Inches(1.6), Inches(11.6), Inches(4.5), font_size=19)

# Slide 11: Conclusion
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
add_title(slide, '10. Conclusion', 'Final summary')
add_bullets(slide, [
    'This project demonstrates how AI can support early awareness and risk assessment in healthcare.',
    'The disease prediction system combines machine learning, web development, and database management into one useful application.',
    'It provides a strong foundation for future innovations in digital healthcare and decision-support systems.',
    'The solution is highly suitable as an educational, prototype, and deployment-ready project for AI in medicine.'
], Inches(0.8), Inches(1.6), Inches(11.5), Inches(4.2), font_size=20)

prs.save(OUTPUT_PATH)
print(f"PPT created successfully: {OUTPUT_PATH}")
